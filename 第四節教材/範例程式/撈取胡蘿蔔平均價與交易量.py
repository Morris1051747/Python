import requests

base_url = 'http://data.coa.gov.tw/Service/OpenData/FromM/FarmTransData.aspx?$skip=0&Market=%E6%BA%AA%E6%B9%96%E9%8E%AE&Crop=%E8%83%A1%E8%98%BF%E8%94%94&StartDate=107.01.01&EndDate=107.12.31'

r = requests.get(base_url)

raw_data = reversed(r.json())

data_box= []
months = []

for s in raw_data:
    if s['作物名稱'] != '休市':
        data= []
        month = s['交易日期'][4:6]
        if month not in months:
            months.append(month)
            
        data.append(month)
        data.append(float(s['平均價']))
        data.append(int(s['交易量']))

        data_box.append(data)

output=[['月份','平均價','交易量']]

for mon in months:
    price = 0
    amount = 0
    time = 0
    for i in data_box:
        if i[0] == mon:
            price = price + i[1]
            amount = amount + i[2]
            time = time + 1
            
    ave_price = round(price / time)
    ave_amount = round(amount / time)

    thing = [mon,ave_price,ave_amount]
    output.append(thing)


import openpyxl
from openpyxl.chart import LineChart, Reference

book = openpyxl.Workbook()

sheet = book.active

for data in output:
    sheet.append(data)
    
price_data = Reference(sheet, min_col=2, min_row=2, max_col=2, max_row=13)
amount_data = Reference(sheet, min_col=3, min_row=2, max_col=3, max_row=13)

chart = LineChart()
chart.add_data(price_data)

chart.title = "各月平均交易價"
chart.y_axis.title = '價錢'
chart.x_axis.title = '月份'
chart.legend = None

sheet.add_chart(chart, "E3")

chart2 = LineChart()
chart2.add_data(amount_data)

chart2.title = "各月平均交易量"
chart2.y_axis.title = '交易量'
chart2.x_axis.title = '月份'
chart2.legend = None

sheet.add_chart(chart2, "O3")

book.save('output.xlsx')
